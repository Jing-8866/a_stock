# -*- coding: utf-8 -*-
import sys, json
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_font(run, name='WenQuanYi Micro Hei', size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color: run.font.color.rgb = RGBColor(*color)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), hexcolor)
    tcPr.append(sh)

def border(cell, color='D9DEE6'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right'):
        e = OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),color)
        borders.append(e)
    tcPr.append(borders)

def hline(doc, color='C0392B', size=18):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(4)
    r = p.add_run('―'*60); set_font(r, size=6, color=(192,57,43))

def add_title(doc, text, size=15):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(8)
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'12'); bottom.set(qn('w:space'),'4'); bottom.set(qn('w:color'),'C0392B')
    pbdr.append(bottom); pPr.append(pbdr)
    r = p.add_run(text); set_font(r, size=size, bold=True, color=(17,17,17))

def add_para(doc, text, size=10.5):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(5); p.paragraph_format.line_spacing = 1.45
    r = p.add_run(text); set_font(r, size=size)

def rows_to_objects(rows):
    if not rows or not rows[0]: return []
    if isinstance(rows[0], dict): return rows
    if all(isinstance(r, list) for r in rows) and len(rows[0]) == 2:
        return [{'项目': r[0], '内容': r[1]} for r in rows]
    return rows

def norm(rows, cols):
    if not rows: return [], []
    data = rows_to_objects(rows)
    if data and isinstance(data[0], dict):
        use = list(data[0].keys())
    else:
        use = list(rows[0])
    if cols: use = [c for c in cols if c in use]
    return use, data

def add_table(doc, rows, cols=None, widths=None):
    if rows and not cols and isinstance(rows[0], dict):
        cols = list(rows[0].keys())
    use, data = norm(rows, cols)
    if not data or not use: return
    label = {'项目': '项目', '内容': '内容'}
    t = doc.add_table(rows=1, cols=len(use)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(use):
        cell = t.rows[0].cells[i]; cell.text = ''; shade(cell, 'C0392B'); border(cell, 'C0392B')
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(str(label.get(c, c))); set_font(r, size=9, bold=True, color=(255,255,255))
    for ri, row in enumerate(data):
        cells = t.add_row().cells
        for i, c in enumerate(use):
            cells[i].text = ''; border(cells[i])
            if ri % 2 == 1: shade(cells[i], 'FAFBFC')
            p = cells[i].paragraphs[0]; r = p.add_run(str(row.get(c, ''))); set_font(r, size=8.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)

def add_box(doc, lines, title='提示', fill='FBF6EC', line='ECDCC0', txt=(107,88,64)):
    box = doc.add_table(rows=1, cols=1); box.style='Table Grid'
    cell = box.cell(0,0); shade(cell, fill.replace('#','')); border(cell, line.replace('#',''))
    for i, ln in enumerate(lines):
        p = cell.paragraphs[0] if i==0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.5
        r = p.add_run(ln); set_font(r, size=9, bold=(i==0 and False), color=txt)

def main():
    data = json.load(open(sys.argv[1], encoding='utf-8'))
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(1.6); sec.left_margin = Cm(1.8); sec.right_margin = Cm(1.8)
    st = doc.styles['Normal']; st.font.name='WenQuanYi Micro Hei'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'WenQuanYi Micro Hei'); st.font.size=Pt(10.5)

    # 封面
    p = doc.add_paragraph(); r = p.add_run('A股选股分析工具 · 自动生成'); set_font(r, size=11, bold=True, color=(255,255,255))
    pPr = p._p.get_or_add_pPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:fill'),'C0392B'); pPr.append(shd)
    p.paragraph_format.space_after = Pt(60)
    p = doc.add_paragraph(); r = p.add_run(data.get('title','股票分析报告')); set_font(r, size=26, bold=True)
    p.paragraph_format.space_after = Pt(10)
    p = doc.add_paragraph(); r = p.add_run('基于腾讯自选股接口实时数据 · 四维评分模型：估值28% / 财务32% / 技术22% / 资金18%'); set_font(r, size=11, color=(100,110,122))
    p.paragraph_format.space_after = Pt(30)
    hline(doc)
    for k, v in data.get('meta', []):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(k+'：'); set_font(r1, size=10.5, bold=True)
        r2 = p.add_run(str(v)); set_font(r2, size=10.5)
    doc.add_page_break()

    # 正文
    for sec2 in data.get('sections', []):
        add_title(doc, sec2.get('title',''))
        if sec2.get('para'):
            for t in sec2['para']: add_para(doc, t)
        if sec2.get('rows'):
            add_table(doc, sec2['rows'], sec2.get('cols'))
        doc.add_paragraph()
    for tb in data.get('tables', []):
        add_title(doc, tb.get('title',''))
        add_table(doc, tb['rows'], tb.get('cols'))
        doc.add_paragraph()
    for ls in data.get('lists', []):
        add_title(doc, ls.get('title',''))
        add_para(doc, ls.get('content',''))
        doc.add_paragraph()

    # 免责声明
    add_box(doc, data.get('notes', ['本报告由 A股选股分析工具 自动生成，为客观数据分析，不构成投资建议。']), '免责声明')

    doc.save(sys.argv[2])

if __name__ == '__main__':
    main()
